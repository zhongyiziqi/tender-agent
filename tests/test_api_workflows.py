from __future__ import annotations

from io import BytesIO

from docx import Document

from tests.conftest import wait_for_task


def make_docx(text: str) -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_heading("测试招标文件", level=1)
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def market_payload(count: int = 3) -> dict:
    products = []
    for index in range(count):
        products.append(
            {
                "product_name": f"产品{index + 1}",
                "vendor_name": f"厂商{index + 1}",
                "price": 100 + index * 10,
                "performance_score": 90 - index,
                "service_score": 85 - index,
                "parameters": {"容量": "100L", "保修": "3年"},
                "source_document_ids": [],
            }
        )
    return {
        "requirement": {
            "title": "测试设备采购",
            "functions": ["稳定运行"],
            "indicators": [{"name": "容量", "operator": ">=", "target": 100, "unit": "L"}],
            "budget_max": 500,
            "service_requirements": ["三年保修"],
        },
        "weights": {"price": 0.4, "performance": 0.4, "service": 0.2},
        "candidate_products": products,
    }


def create_and_wait(client, endpoint: str, payload: dict) -> tuple[str, dict]:
    response = client.post(endpoint, json=payload)
    assert response.status_code == 202, response.text
    task_id = response.json()["id"]
    task = wait_for_task(client, task_id)
    return task_id, task


def test_health_and_document_upload(client):
    health = client.get("/api/v1/health")
    assert health.status_code == 200
    assert health.json()["llm_mode"] == "mock"

    response = client.post(
        "/api/v1/documents",
        files={"file": ("sample.docx", make_docx("采购文件正文"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201, response.text
    document_id = response.json()["id"]
    assert client.get(f"/api/v1/documents/{document_id}").status_code == 200

    invalid = client.post(
        "/api/v1/documents", files={"file": ("bad.txt", b"text", "text/plain")}
    )
    assert invalid.status_code == 422


def test_three_agent_workflow_and_artifacts(client):
    market_id, market_task = create_and_wait(client, "/api/v1/tasks/market-research", market_payload())
    assert market_task["status"] == "completed"
    market_result = client.get(f"/api/v1/tasks/{market_id}/result")
    assert market_result.status_code == 200
    assert len(market_result.json()["result"]["top_three"]) == 3
    assert {item["kind"] for item in market_result.json()["artifacts"]} == {"json", "docx", "pdf"}

    tender_id, tender_task = create_and_wait(
        client,
        "/api/v1/tasks/tender-generation",
        {"market_research_task_id": market_id},
    )
    assert tender_task["status"] == "completed"
    tender_result = client.get(f"/api/v1/tasks/{tender_id}/result").json()
    assert tender_result["result"]["vendor_coverage_count"] == 3

    review_id, review_task = create_and_wait(
        client,
        "/api/v1/tasks/compliance-review",
        {"generation_task_id": tender_id},
    )
    assert review_task["status"] == "completed"
    review_result = client.get(f"/api/v1/tasks/{review_id}/result").json()
    assert review_result["result"]["risk_counts"]["high"] == 0
    for artifact in review_result["artifacts"]:
        download = client.get(artifact["download_url"])
        assert download.status_code == 200
        assert download.content


def test_insufficient_coverage_and_uploaded_document_review(client):
    market_id, market_task = create_and_wait(
        client, "/api/v1/tasks/market-research", market_payload(count=2)
    )
    assert market_task["status"] == "completed"
    _, tender_task = create_and_wait(
        client,
        "/api/v1/tasks/tender-generation",
        {"market_research_task_id": market_id},
    )
    assert tender_task["status"] == "needs_review"

    upload = client.post(
        "/api/v1/documents",
        files={"file": ("risky.docx", make_docx("本项目指定品牌，且仅限本地注册供应商。"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    document_id = upload.json()["id"]
    review_id, review_task = create_and_wait(
        client,
        "/api/v1/tasks/compliance-review",
        {"document_id": document_id, "vendor_coverage_count": 2},
    )
    assert review_task["status"] == "needs_review"
    result = client.get(f"/api/v1/tasks/{review_id}/result").json()["result"]
    assert result["risk_counts"]["high"] >= 2
    assert any(item["category"] == "厂商覆盖性" for item in result["findings"])

