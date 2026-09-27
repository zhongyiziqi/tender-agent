from __future__ import annotations

import hashlib
from pathlib import Path
from uuid import uuid4

from docx import Document as WordDocument
from fastapi import UploadFile
from pypdf import PdfReader

from app.core.config import Settings
from app.models import Document
from app.repositories.repository import Repository


ALLOWED_SUFFIXES = {".docx", ".pdf"}
CONTENT_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
}


def _chunk_text(text: str, size: int = 1200, overlap: int = 150) -> list[str]:
    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(len(normalized), start + size)
        chunks.append(normalized[start:end])
        if end == len(normalized):
            break
        start = end - overlap
    return chunks


def parse_file(path: Path) -> list[dict[str, object]]:
    if path.suffix.lower() == ".docx":
        document = WordDocument(path)
        chunks: list[dict[str, object]] = []
        section = "正文"
        buffer: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                if buffer:
                    for part in _chunk_text("\n".join(buffer)):
                        chunks.append({"section": section, "page_number": None, "content": part})
                    buffer.clear()
                section = text
            else:
                buffer.append(text)
        if buffer:
            for part in _chunk_text("\n".join(buffer)):
                chunks.append({"section": section, "page_number": None, "content": part})
        for table_index, table in enumerate(document.tables, start=1):
            table_text = "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells) for row in table.rows
            )
            for part in _chunk_text(table_text):
                chunks.append(
                    {"section": f"表格 {table_index}", "page_number": None, "content": part}
                )
        return chunks

    reader = PdfReader(str(path))
    chunks = []
    for page_number, page in enumerate(reader.pages, start=1):
        for part in _chunk_text(page.extract_text() or ""):
            chunks.append(
                {"section": f"第 {page_number} 页", "page_number": page_number, "content": part}
            )
    return chunks


class DocumentService:
    def __init__(self, settings: Settings, repository: Repository) -> None:
        self.settings = settings
        self.repository = repository

    async def save_upload(self, upload: UploadFile) -> Document:
        suffix = Path(upload.filename or "").suffix.lower()
        if suffix not in ALLOWED_SUFFIXES:
            raise ValueError("仅支持 DOCX 和 PDF 文件")
        content = await upload.read(self.settings.max_upload_mb * 1024 * 1024 + 1)
        if len(content) > self.settings.max_upload_mb * 1024 * 1024:
            raise ValueError(f"文件不能超过 {self.settings.max_upload_mb} MB")
        if not content:
            raise ValueError("文件内容为空")
        target = self.settings.upload_dir / f"{uuid4()}{suffix}"
        target.write_bytes(content)
        try:
            chunks = parse_file(target)
            if not chunks:
                raise ValueError("未能从文件中提取有效文本")
            return self.repository.create_document(
                filename=Path(upload.filename or target.name).name,
                content_type=CONTENT_TYPES[suffix],
                path=target,
                sha256=hashlib.sha256(content).hexdigest(),
                size_bytes=len(content),
                chunks=chunks,
            )
        except Exception:
            target.unlink(missing_ok=True)
            raise

