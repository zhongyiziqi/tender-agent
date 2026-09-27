from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.pdfmetrics import registerFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _lines(value: Any, prefix: str = "") -> list[str]:
    if isinstance(value, dict):
        result: list[str] = []
        for key, item in value.items():
            label = f"{prefix}{key}"
            if isinstance(item, (dict, list)):
                result.append(label)
                result.extend(_lines(item, prefix="  "))
            else:
                result.append(f"{label}: {item}")
        return result
    if isinstance(value, list):
        result = []
        for index, item in enumerate(value, start=1):
            if isinstance(item, (dict, list)):
                result.append(f"{prefix}{index}.")
                result.extend(_lines(item, prefix="  "))
            else:
                result.append(f"{prefix}{index}. {item}")
        return result
    return [f"{prefix}{value}"]


class ArtifactExporter:
    def __init__(self, artifact_dir: Path) -> None:
        self.artifact_dir = artifact_dir

    def export(self, task_id: str, title: str, result: dict[str, Any]) -> dict[str, Path]:
        task_dir = self.artifact_dir / task_id
        task_dir.mkdir(parents=True, exist_ok=True)
        json_path = task_dir / "result.json"
        docx_path = task_dir / "report.docx"
        pdf_path = task_dir / "report.pdf"
        json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        self._write_docx(docx_path, title, result)
        self._write_pdf(pdf_path, title, result)
        return {"json": json_path, "docx": docx_path, "pdf": pdf_path}

    @staticmethod
    def _write_docx(path: Path, title: str, result: dict[str, Any]) -> None:
        document = Document()
        styles = document.styles
        styles["Normal"].font.name = "Microsoft YaHei"
        styles["Normal"].font.size = Pt(10.5)
        document.add_heading(title, level=0)
        document.add_paragraph("本文件为智能体生成草案，须经采购、法务或相关责任人复核确认。")
        for line in _lines(result):
            document.add_paragraph(line)
        document.save(path)

    @staticmethod
    def _write_pdf(path: Path, title: str, result: dict[str, Any]) -> None:
        registerFont(UnicodeCIDFont("STSong-Light"))
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ChineseTitle",
            parent=styles["Title"],
            fontName="STSong-Light",
            fontSize=18,
            leading=24,
            alignment=TA_CENTER,
        )
        body_style = ParagraphStyle(
            "ChineseBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=9, leading=14
        )
        story = [Paragraph(title, title_style), Spacer(1, 12)]
        story.append(Paragraph("本文件为智能体生成草案，须经人工复核确认。", body_style))
        story.append(Spacer(1, 8))
        for line in _lines(result):
            safe = str(line).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe.replace("  ", "&nbsp;&nbsp;"), body_style))
            story.append(Spacer(1, 3))
        SimpleDocTemplate(str(path)).build(story)

